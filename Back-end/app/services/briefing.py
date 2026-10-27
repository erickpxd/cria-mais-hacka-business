"""Lê o briefing (PDF, DOCX ou TXT) e organiza os campos, indicando a página onde cada um foi achado.

Heurística: procura rótulos em MAIÚSCULAS (PRODUTO, PREÇO..., CANAIS..., INFORMAÇÕES OBRIGATÓRIAS...)
ou no formato "Rótulo: valor". Briefings muito diferentes disso vão exigir ajuste nas regras
(ou uma etapa de extração mais inteligente), e o usuário SEMPRE confere antes de criar o projeto."""
import io
import re
import unicodedata
from datetime import date
from pathlib import Path

from app.config import settings
from app.services.erros import ErroNegocio
from app.services.formatos import detectar_canais

ROTULOS = {
    "produto": ("PRODUTO",),
    "preco": ("PRECO",),
    "canais": ("CANAIS", "CANAL"),
    "informacoes_obrigatorias": ("INFORMACOES OBRIGATORIAS",),
    "prazo": ("PRAZO", "DATAS", "DATA"),
    "objetivo": ("OBJETIVO",),
    "orientacoes": ("ORIENTACOES",),
}
RUIDO = re.compile(r"(material demonstrativo|p[aá]gina\s*\d+)", re.I)
IGNORAR_TITULO = re.compile(r"(briefing|dados fict|demonstra)", re.I)
RE_PRECO = re.compile(r"(?<![\d.,])(\d{1,3}(?:\.\d{3})*|\d+),(\d{2})(?!\d)")
RE_DATA = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{2,4})")


def _norm(s: str) -> str:
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().upper().strip()


def ler_paginas(nome: str, conteudo: bytes) -> tuple[list[str], bool]:
    """Devolve (textos por página, paginado). DOCX e TXT não têm página confiável."""
    ext = Path(nome).suffix.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader

            paginas_pdf = PdfReader(io.BytesIO(conteudo)).pages
            if len(paginas_pdf) > settings.max_briefing_pages:
                raise ErroNegocio(413, f"O PDF excede {settings.max_briefing_pages} páginas.")
            paginas = [(p.extract_text() or "") for p in paginas_pdf]
            if any(not pagina.strip() for pagina in paginas):
                paginas = _ocr_paginas_pdf(conteudo, paginas)
            return paginas, True
        if ext == ".docx":
            import docx

            doc = docx.Document(io.BytesIO(conteudo))
            linhas = [p.text for p in doc.paragraphs]
            for tabela in doc.tables:
                for linha in tabela.rows:
                    linhas.append(" | ".join(celula.text.strip() for celula in linha.cells))
            return ["\n".join(linhas)], False
        if ext == ".txt":
            try:
                return [conteudo.decode("utf-8-sig")], False
            except UnicodeDecodeError:
                return [conteudo.decode("latin-1")], False
    except ErroNegocio:
        raise
    except Exception as exc:
        raise ErroNegocio(422, "Não foi possível ler o documento. Confira se o arquivo não está corrompido.") from exc
    raise ErroNegocio(415, "Envie um arquivo PDF, DOCX ou TXT.")


def _ocr_paginas_pdf(conteudo: bytes, paginas: list[str]) -> list[str]:
    """Aplica OCR somente às páginas sem texto selecionável e preserva a paginação."""
    import cv2
    import numpy as np
    import pypdfium2 as pdfium

    from app.services import ocr

    documento = pdfium.PdfDocument(conteudo)
    motor = ocr.obter_motor()
    resultado = list(paginas)
    for indice, texto in enumerate(resultado):
        if texto.strip():
            continue
        pil = documento[indice].render(scale=2).to_pil().convert("RGB")
        imagem = cv2.cvtColor(np.asarray(pil), cv2.COLOR_RGB2BGR)
        linhas = sorted(motor.ler(imagem), key=lambda item: (item.y, item.x))
        resultado[indice] = "\n".join(
            linha.texto for linha in linhas if linha.confianca >= settings.ocr_confianca_min
        )
    return resultado


def _detectar_rotulo(linha: str) -> tuple[str | None, str]:
    n = _norm(linha)
    for chave, prefixos in ROTULOS.items():
        for p in prefixos:
            if not n.startswith(p):
                continue
            if ":" in linha[:60]:  # "Rótulo: valor"
                return chave, linha.split(":", 1)[1].strip()
            if linha.strip() == linha.strip().upper():  # linha de rótulo em maiúsculas
                return chave, ""
    return None, ""


def parse_preco(texto: str | None) -> int | None:
    """'R$ 17,90' / '17,90' / '1.299,00' -> centavos."""
    if not texto:
        return None
    m = RE_PRECO.search(texto)
    if not m:
        return None
    return int(m.group(1).replace(".", "")) * 100 + int(m.group(2))


def formatar_preco(centavos: int) -> str:
    reais, cent = divmod(centavos, 100)
    return f"{reais:,}".replace(",", ".") + f",{cent:02d}"


def _parse_data(texto: str) -> date | None:
    m = RE_DATA.search(texto)
    if not m:
        return None
    d, mth, a = int(m.group(1)), int(m.group(2)), int(m.group(3))
    a = a + 2000 if a < 100 else a
    try:
        return date(a, mth, d)
    except ValueError:
        return None


def extrair_campos(paginas: list[str], paginado: bool) -> dict:
    campos = {k: {"valor": None, "pagina": None, "encontrado": False, "estado": "not_found"} for k in
              ("nome_campanha", "produto", "preco", "canais", "prazo",
               "informacoes_obrigatorias", "objetivo", "orientacoes")}
    blocos: dict[str, list[str]] = {}
    pagina_do_bloco: dict[str, int | None] = {}
    rotulo_atual = None

    for i, texto in enumerate(paginas):
        pagina = i + 1 if paginado else None
        for linha in (l.strip() for l in texto.splitlines()):
            if not linha:
                continue
            chave, resto = _detectar_rotulo(linha)
            if chave:
                rotulo_atual = chave
                blocos.setdefault(chave, [])
                pagina_do_bloco.setdefault(chave, pagina)
                if resto:
                    blocos[chave].append(resto)
                continue
            if RUIDO.search(linha):
                continue
            if rotulo_atual is None:
                if not campos["nome_campanha"]["encontrado"] and not IGNORAR_TITULO.search(linha):
                    campos["nome_campanha"] = {"valor": linha, "pagina": pagina,
                                                "encontrado": True, "estado": "found",
                                                "trecho": linha}
                continue
            blocos[rotulo_atual].append(linha)

    def preencher(chave_campo: str, chave_bloco: str, valor):
        if valor:
            campos[chave_campo] = {"valor": valor, "pagina": pagina_do_bloco.get(chave_bloco),
                                   "encontrado": True, "estado": "found",
                                   "trecho": " ".join(blocos.get(chave_bloco, [])).strip()}

    for chave in ("produto", "informacoes_obrigatorias", "objetivo", "orientacoes"):
        preencher(chave, chave, " ".join(blocos.get(chave, [])).strip())

    texto_preco = " ".join(blocos.get("preco", []))
    precos = list(dict.fromkeys(f"{m.group(1)},{m.group(2)}" for m in RE_PRECO.finditer(texto_preco)))
    if len(precos) == 1:
        preencher("preco", "preco", precos[0])
    elif len(precos) > 1:
        campos["preco"] = {"valor": None, "pagina": pagina_do_bloco.get("preco"),
                           "encontrado": False, "estado": "conflict", "candidatos": precos,
                           "trecho": texto_preco}

    preencher("canais", "canais", detectar_canais(" ".join(blocos.get("canais", []))))

    data = _parse_data(" ".join(blocos.get("prazo", [])))
    preencher("prazo", "prazo", data.isoformat() if data else None)
    return campos


def avisos_do_briefing(campos: dict, paginas: list[str]) -> list[str]:
    avisos = []
    if not any(p.strip() for p in paginas):
        avisos.append("O documento não tem texto selecionável (pode ser um arquivo escaneado). "
                      "Preencha as informações manualmente.")
    if not campos["prazo"]["encontrado"]:
        avisos.append("Prazo não informado. Você pode definir depois, sem alterar o documento original.")
    if campos["preco"].get("estado") == "conflict":
        avisos.append("Mais de um preço foi encontrado. Confirme manualmente qual valor pertence ao produto.")
    return avisos
