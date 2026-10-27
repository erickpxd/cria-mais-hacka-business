import io

import pytest

from app.domain import StatusArte, transicionar
from app.services.briefing import extrair_campos, formatar_preco, ler_paginas, parse_preco
from app.services.erros import ErroNegocio


def test_preco_brl_em_centavos():
    assert parse_preco("R$ 1.299,90") == 129990
    assert formatar_preco(129990) == "1.299,90"


def test_preco_ausente_nao_e_inventado():
    assert parse_preco(None) is None
    assert parse_preco("sem preço") is None


def test_fluxo_de_aprovacao_valido():
    assert transicionar("em_producao", StatusArte.AGUARDANDO_APROVACAO) == "aguardando_aprovacao"
    assert transicionar("aguardando_aprovacao", StatusArte.APROVADA) == "aprovada"


def test_aprovacao_direta_invalida():
    with pytest.raises(ErroNegocio) as error:
        transicionar("em_producao", StatusArte.APROVADA)
    assert error.value.status == 409


def test_briefing_nao_escolhe_preco_ambiguo():
    campos = extrair_campos(["CAMPANHA TESTE\nPRODUTO: Shampoo\nPREÇO: R$ 17,90 ou R$ 19,90"], False)
    assert campos["preco"]["estado"] == "conflict"
    assert campos["preco"]["valor"] is None
    assert campos["preco"]["candidatos"] == ["17,90", "19,90"]


def test_briefing_nao_inventa_prazo():
    campos = extrair_campos(["CAMPANHA TESTE\nPRODUTO: Shampoo"], False)
    assert campos["prazo"]["estado"] == "not_found"
    assert campos["prazo"]["valor"] is None


def test_docx_le_paragrafos_e_tabelas():
    import docx

    documento = docx.Document()
    documento.add_paragraph("CAMPANHA TESTE")
    tabela = documento.add_table(rows=2, cols=2)
    tabela.cell(0, 0).text = "PRODUTO:"
    tabela.cell(0, 1).text = "Shampoo 400 ml"
    tabela.cell(1, 0).text = "PREÇO:"
    tabela.cell(1, 1).text = "R$ 17,90"
    arquivo = io.BytesIO()
    documento.save(arquivo)

    paginas, paginado = ler_paginas("briefing.docx", arquivo.getvalue())
    assert paginado is False
    assert "Shampoo 400 ml" in paginas[0]
    assert "R$ 17,90" in paginas[0]


def test_pdf_escaneado_usa_ocr(monkeypatch):
    from PIL import Image

    from app.services import ocr

    class MotorPDF:
        def ler(self, _imagem):
            return [ocr.LinhaTexto("PRODUTO: Shampoo 400 ml", 10, 10, 200, 20, 0.99)]

    monkeypatch.setattr(ocr, "obter_motor", lambda: MotorPDF())
    arquivo = io.BytesIO()
    Image.new("RGB", (300, 200), "white").save(arquivo, format="PDF")
    paginas, paginado = ler_paginas("escaneado.pdf", arquivo.getvalue())
    assert paginado is True
    assert paginas == ["PRODUTO: Shampoo 400 ml"]
